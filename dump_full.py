import openpyxl

F = r"B2C Logistics Commercials (2).xlsx"
OUT = "dump_full.txt"
wb = openpyxl.load_workbook(F, read_only=True, data_only=True)
lines = []
lines.append("SHEETS: " + str(wb.sheetnames))

for name in ["Delhivery", "Bluedart", "DTDC", "Ekart", "Shadowfax", "Velocity Express", "Amazon", "Elastic Run"]:
    ws = wb[name]
    lines.append("\n" + "=" * 110)
    lines.append(f"### SHEET: {name}  max_row={ws.max_row} max_col={ws.max_column}")
    lines.append("=" * 110)
    for r in ws.iter_rows():
        cells = []
        for c in r:
            if c.value is not None:
                v = c.value
                if isinstance(v, float) and v == int(v):
                    v = int(v)
                cells.append(f"[{c.coordinate}]{v}")
        if cells:
            lines.append(" | ".join(cells))

with open(OUT, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("written", OUT, len("\n".join(lines)), "chars")