import openpyxl

F = r"B2C Logistics Commercials (2).xlsx"
OUT = "dump_support.txt"
wb = openpyxl.load_workbook(F, read_only=True, data_only=True)
lines = []
for name in ["WH Details", "Carrier List", "Billing Zone Mapping", "Actual Zone Name"]:
    ws = wb[name]
    lines.append(f"### SHEET: {name}  max_row={ws.max_row} max_col={ws.max_column}")
    lines.append("-" * 110)
    for r in ws.iter_rows():
        cells = []
        for c in r:
            if c.value is not None:
                v = c.value
                if isinstance(v, float) and v == int(v):
                    v = int(v)
                cells.append(f"[{c.coordinate}]{v}")
        if cells:
            lines.append(" | ".join(cells))
    lines.append("")
with open(OUT, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("written", OUT, len("\n".join(lines)), "chars")