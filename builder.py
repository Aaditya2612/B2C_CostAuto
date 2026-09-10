"""Builds the app's data artifacts from the source Excel workbooks.

Artifacts produced in ./data:
  - lanes.db          SQLite index of the Billing Zone Master "Main" lane sheet
  - b2c_data.json     warehouses, active carriers, zone mapping tables, price cards

Run:  python builder.py
"""
from __future__ import annotations

import json
import os
import sqlite3

import openpyxl
import pyxlsb

ROOT = os.path.dirname(os.path.abspath(__file__))
B2C_FILE = os.path.join(ROOT, "B2C Logistics Commercials (2).xlsx")
ZONE_FILE = os.path.join(ROOT, "Billing Zone Master Updated.xlsb")
DATA_DIR = os.path.join(ROOT, "data")
LANES_DB = os.path.join(DATA_DIR, "lanes.db")
B2C_JSON = os.path.join(DATA_DIR, "b2c_data.json")

# Carriers that have a commercial sheet in the B2C workbook are the ONLY active
# ones. Everything else in "Carrier List" is ignored.
ACTIVE_CARRIERS = [
    {"id": "delhivery", "name": "Delhivery", "master_col": "Delhivery", "sheet": "Delhivery"},
    {"id": "bluedart", "name": "Bluedart (Dart Plus)", "master_col": "Bluedart Plus", "sheet": "Bluedart"},
    {"id": "dtdc", "name": "DTDC", "master_col": "DTDC", "sheet": "DTDC"},
    {"id": "ekart", "name": "Ekart", "master_col": "Ekart", "sheet": "Ekart"},
    {"id": "shadowfax", "name": "Shadowfax", "master_col": "Shadowfax", "sheet": "Shadowfax"},
    {"id": "velocity", "name": "Velocity Express", "master_col": "Velocity Express", "sheet": "Velocity Express"},
    {"id": "amazon", "name": "Amazon", "master_col": "Amazon", "sheet": "Amazon"},
    {"id": "elastic", "name": "Elastic Run", "master_col": None, "sheet": "Elastic Run"},  # no lane zone column
]

# Sentinel values in the Zone Master that mean "carrier does not serve this lane".
NA_SENTINELS = {"NA", "N/A", "", "0x2a", "*", "\ufeffNA"}


def norm(v):
    """Normalise a zone cell value. Returns None for not-served sentinels."""
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    low = s.upper().replace("\u00a0", " ").replace("\ufeff", "")
    if low in NA_SENTINELS or "0x2a" in low.lower():
        return None
    return s


def extract_wh_details(wb):
    ws = wb["WH Details"]
    rows = {}
    for r in ws.iter_rows(min_row=4, values_only=True):
        if r[0] is None and r[3] is None:
            continue
        code, state, gstn, whid = r[0], r[1], r[2], r[3]
        addr, origin_pin = r[4], r[5]
        if whid is None or str(whid).strip() == "":
            continue
        rows[int(whid)] = {
            "whid": int(whid),
            "code": (str(code).strip() if code is not None else ""),
            "state": (str(state).strip() if state is not None else ""),
            "gstn": (str(gstn).strip() if gstn is not None else ""),
            "address": (str(addr).strip() if addr is not None else ""),
            "origin_pin": int(origin_pin) if origin_pin is not None else None,
        }
    return rows


def extract_zone_maps(wb):
    """Billing Zone Mapping (carrier -> carrier_zone -> system zone) and
    Actual Zone Name (actual_zone -> final zone)."""
    billing = {}     # carrier name normalized -> {carrier_zone: system_zone}
    for r in wb["Billing Zone Mapping"].iter_rows(min_row=2, values_only=True):
        system_zone, carrier_name, carrier_zone = r[1], r[2], r[3]
        if not carrier_name or not carrier_zone:
            continue
        key = str(carrier_name).strip().lower()
        billing.setdefault(key, {})[str(carrier_zone).strip()] = str(system_zone).strip()

    actual = {}      # actual_zone -> final zone
    for r in wb["Actual Zone Name"].iter_rows(min_row=2, values_only=True):
        az, z = r[0], r[1]
        if az is None:
            continue
        actual[str(az).strip()] = str(z).strip()
    return billing, actual


def extract_lanes():
    os.makedirs(DATA_DIR, exist_ok=True)
    if os.path.exists(LANES_DB):
        os.remove(LANES_DB)
    con = sqlite3.connect(LANES_DB)
    cur = con.cursor()
    cur.execute(
        """CREATE TABLE lanes (
            whid INTEGER, origin_pin INTEGER, pin INTEGER, city TEXT, state TEXT, lane TEXT,
            ecom TEXT, delhivery TEXT, delhivery_mm TEXT, shadowfax TEXT, bluedart_plus TEXT,
            dtdc TEXT, velocity TEXT, delhivery_ndd TEXT, ekart TEXT, amazon TEXT
        )"""
    )
    cur.execute("CREATE INDEX idx_lanes ON lanes(whid, pin)")

    mapping = [
        ("ecom", "Ecom Express"),
        ("delhivery", "Delhivery"),
        ("delhivery_mm", "Delhivery_MM"),
        ("shadowfax", "Shadowfax"),
        ("bluedart_plus", "Bluedart Plus"),
        ("dtdc", "DTDC"),
        ("velocity", "Velocity Express"),
        ("delhivery_ndd", "Delhivery_NDD"),
        ("ekart", "Ekart"),
        ("amazon", "Amazon"),
    ]

    cols = None
    idx = {}
    batch = []
    n = 0
    zone_values = {}   # master col -> set of distinct zone values encountered
    with pyxlsb.open_workbook(ZONE_FILE) as wb:
        with wb.get_sheet("Main") as sh:
            for row in sh.rows():
                vals = [c.v for c in row]
                if cols is None:
                    cols = [str(c).strip() for c in vals]
                    idx = {name: i for i, name in enumerate(cols)}
                    continue
                def gi(name):
                    v = vals[idx[name]]
                    return int(round(float(v))) if v is not None else None

                def g(name):
                    return norm(vals[idx[name]])

                whid = gi("WHID")
                pin = gi("Destination")
                if whid is None or pin is None:
                    continue
                city = vals[idx["City"]]
                state = vals[idx["State"]]
                rowvals = {
                    "whid": whid,
                    "origin_pin": gi("Origin"),
                    "pin": pin,
                    "city": str(city).strip() if city is not None else "",
                    "state": str(state).strip() if state is not None else "",
                    "lane": str(vals[idx["Lane"]]).strip() if vals[idx["Lane"]] is not None else f"{whid}_{pin}",
                }
                for dom, src in mapping:
                    rowvals[dom] = g(src)
                    if src not in zone_values:
                        zone_values[src] = set()
                    if rowvals[dom]:
                        zone_values[src].add(rowvals[dom])
                batch.append((
                    rowvals["whid"], rowvals["origin_pin"], rowvals["pin"],
                    rowvals["city"], rowvals["state"], rowvals["lane"],
                    rowvals["ecom"], rowvals["delhivery"], rowvals["delhivery_mm"],
                    rowvals["shadowfax"], rowvals["bluedart_plus"], rowvals["dtdc"],
                    rowvals["velocity"], rowvals["delhivery_ndd"], rowvals["ekart"],
                    rowvals["amazon"],
                ))
                n += 1
                if len(batch) >= 5000:
                    cur.executemany(
                        """INSERT INTO lanes
                           (whid, origin_pin, pin, city, state, lane, ecom, delhivery,
                            delhivery_mm, shadowfax, bluedart_plus, dtdc, velocity,
                            delhivery_ndd, ekart, amazon)
                           VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                        batch,
                    )
                    batch = []
            if batch:
                cur.executemany(
                    """INSERT INTO lanes
                       (whid, origin_pin, pin, city, state, lane, ecom, delhivery,
                        delhivery_mm, shadowfax, bluedart_plus, dtdc, velocity,
                        delhivery_ndd, ekart, amazon)
                       VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    batch,
                )
    con.commit()
    con.close()
    return n, zone_values


# ---------------------------------------------------------------------------
# Price cards authored from the B2C Logistics Commercials workbook.
# Each carrier implements its own "minute rules" from its sheet.
# ---------------------------------------------------------------------------

PRICE_CARDS = {
    "delhivery": {
        "label": "Delhivery Surface (W.E.F 01-02-2024)",
        "zones": ["A", "B", "C1", "C2", "D1", "D2", "E", "F", "B_SPL"],
        "base_forward": {  # first 0-250 g slab, per zone (A/B are vars below)
            "A": "X", "B": "Y", "C1": 55.0, "C2": 55.0,
            "D1": 65.0, "D2": 65.0, "E": 85.0, "F": 92.0,
        },
        "xy": {  # monthly volume slabs -> A (X) / B (Y) forward rates
            # <1 Lakh = 'Existing rate card, Table B'; X/Y are NOT printed anywhere
            # in the workbook, so they are intentionally unavailable (no made-up numbers).
            "<1 Lakh": {"X": None, "Y": None},
            "1-3 Lakh": {"X": 35.0, "Y": 42.0},
            ">=3 Lakh": {"X": 31.0, "Y": 38.5},
        },
        "b_spl": 42.0,        # Guwahati (B_SPL) special: Rs 42 / shipment
        "dto": {              # DTO movement: first 500 g then per additional 500 g
            "base_500": {"A": 44.0, "B": 44.0, "C1": 55.7, "C2": 55.7,
                         "D1": 64.7, "D2": 64.7, "E": 70.0, "F": 74.0},
            "addl_500": {"A": 40.0, "B": 40.0, "C1": 53.7, "C2": 53.7,
                         "D1": 62.7, "D2": 62.7, "E": 68.0, "F": 71.0},
        },
        "rto": 0.0,
        "cod": 0.0,
        "movements": ["Fwd", "RTO", "DTO"],
    },
    "bluedart": {
        "label": "Dart Plus Rates (valid till 31-Dec-2025)",
        "zones": ["Same City", "Same Region", "Metro", "Rest of India", "NE/JNK"],
        "master_to_rate": {  # zone value in Zone Master -> rate-card column
            "Intra City": "Same City",
            "Intra city": "Same City",
            "INTRA CITY": "Same City",
            "Intra Region": "Same Region",
            "Intra region": "Same Region",
            "Metro": "Metro",
            "ROI": "Rest of India",
            "Rest of India": "Rest of India",
            "NE/JNK": "NE/JNK",
            "NE J&K": "NE/JNK",
        },
        "first_1kg": {"Same City": 21.5, "Same Region": 30.0, "Metro": 40.0,
                      "Rest of India": 43.0, "NE/JNK": 46.0},
        "addl_500": {"Same City": 5.0, "Same Region": 7.0, "Metro": 13.0,
                     "Rest of India": 18.0, "NE/JNK": 21.0},
        "min_weight_kg": 1.0,
        "rto_factor": 0.7,
        "vol_discount": {  # monthly shipments -> discount fraction
            "Up to 1,50,000": 0.0,
            "1,50,001 to 3,00,000": 0.025,
            "3,00,001 to 4,50,000": 0.05,
            "4,50,001 to 7,50,000": 0.075,
            "7,50,001 and above": 0.1,
        },
        "movements": ["Fwd", "RTO"],
    },
    "dtdc": {
        "label": "DTDC (w.e.f 21-May-2025)",
        "zones": ["A-Intracity", "Region", "B-Within Zone", "C-Metro", "A-ROI", "ROI-B", "E-Special Zone"],
        "first_5000g": 34.9,
        "addl_1000g": 0.0,
        "movements": ["Fwd"],   # card quotes forward only (no RTO/DTO listed)
    },
    "ekart": {
        "label": "Ekart Base rates (0-5 kg + 500 g thereafter)",
        "vol_discount": {  # monthly picked volume -> (0-5kg rate, addl 500g rate); scheme "From Jul'26 onwards"
            "Base rates (Rs 41)": {"first_5kg": 41.0, "addl_500": 7.0},
            ">=2,00,000 /mo (Rs 38)": {"first_5kg": 38.0, "addl_500": 7.0},
            ">=4,00,000 /mo (Rs 37)": {"first_5kg": 37.0, "addl_500": 7.0},
        },
        "note": ("Base rates inclusive of forward, COD & RTO charges. Volume discounts per the sheet: "
                 "May'26-Jun'26  >=1,50,000 => Rs 38, >=3,00,000 => Rs 37; "
                 "from Jul'26 onwards  >=2,00,000 => Rs 38, >=4,00,000 => Rs 37."),
        "movements": ["Fwd", "RTO"],  # base inclusive of forward + RTO (no separate DTO card)
    },
    "shadowfax": {
        "label": "Shadowfax (W.E.F 01-06-2025)",
        "zones": ["Intracity", "Within Zone", "Intrastate", "Metro", "ROI", "Special Zone"],
        "rates": {  # zone category -> (low vol <4L, high vol >=4L)
            "Intracity": (33.0, 33.0),          # A
            "Intrastate": (35.5, 33.5),         # B1 Within state / Intrastate
            "Within Zone": (38.5, 36.5),        # B2 Within Zone / Regional
            "Within zone": (38.5, 36.5),
            "Regional": (38.5, 36.5),           # master emits 'Regional' for category B2
            "Metro": (42.5, 42.5),             # C Metro/ROI/SZ
            "ROI": (42.5, 42.5),
            "Special Zone": (42.5, 42.5),
            "Only LM": (31.0, 31.0),           # G Kerala (Only LM)
            "Kerala (Only LM)": (31.0, 31.0),
        },
        "kerala_lm": 31.0,
        "movements": ["Fwd"],
        "note": "Sheet also lists a separate 'RVP with QC' product (reverse pickup) with its own CPS; not a forward movement.",
    },
    "velocity": {
        "label": "Velocity Express (incl carrying cost)",
        "zones": ["Local", "Intrastate", "Regional", "Intra Region", "Only LM"],
        "rates": {
            "Only LM": 29.0,
            "Local": 31.0,
            "Intrastate": 36.0,
            "Intra Region": 40.0,
            "Regional": 40.0,
        },
        "movements": ["Fwd"],
    },
    "amazon": {
        "label": "Amazon CPS",
        "zones": ["Local", "Metro", "Regional", "National", "Remote"],
        "cps": {  # monthly shipment volume slab -> CPS (col A of the sheet)
            "<2 Lakh": 40.0,
            "2-3 Lakh": 36.0,
            "3-4 Lakh": 35.2,
            "4-5 Lakh": 34.8,
            ">5 Lakh": 34.4,
        },
        "disc": {  # Cumulative Disc (col B of the sheet)
            "<2 Lakh": 0.0,
            "2-3 Lakh": 0.10,
            "3-4 Lakh": 0.12,
            "4-5 Lakh": 0.13,
            ">5 Lakh": 0.14,
        },
        "base_kg_inclusive": 2.0,
        "addl_per_kg": 15.0,   # discount slabs apply on additional per-kg rates too
        "movements": ["Fwd"],
        "note": "Additional INR 15/Kg billed per full kg above 2 kg.",
    },
    "elastic": {
        "label": "Elastic Run",
        "vol_rates": {  # avg monthly volume -> rate per shipment (From 01-02-2025)
            "<4000": 33.0,
            "4000-4500": 32.0,
            ">4500": 31.0,
        },
        "sdd_rate": 37.0,                 # SDD Rs 37 per delivered order (From 01-07-2025)
        "wh_sdd_37_ids": [2, 4, 10, 12, 28],
        "ndd_regional": 40.0,              # NDD - Regional (From 20-02-2026), per delivered order
        "movements": ["Fwd"],             # NDD: "No return freight, bill on delivered shipments only"
    },
}

MASTER_COL_BY_ID = {c["id"]: c["master_col"] for c in ACTIVE_CARRIERS}
# Mapping for the zone-chain display (carrier zone -> final canonical zone).
# Uses system-zone from Billing Zone Mapping then Actual Zone Name.
# For zones not present there we fall back to the Actual Zone Name directly.


def main():
    if not os.path.exists(B2C_FILE) or not os.path.exists(ZONE_FILE):
        raise SystemExit(f"Missing source files:\n  {B2C_FILE}\n  {ZONE_FILE}")
    os.makedirs(DATA_DIR, exist_ok=True)

    wb = openpyxl.load_workbook(B2C_FILE, read_only=True, data_only=True)
    whs = extract_wh_details(wb)
    billing, actual = extract_zone_maps(wb)

    print("Building lanes.db from Zone Master 'Main' (this takes a bit)...")
    n_lanes, zone_values = extract_lanes()
    print(f"  {n_lanes} lanes written")

    # Distinct zone values seen per master column -> used for the UI dropdowns
    master_zones = {c["id"]: sorted(zone_values.get(c["master_col"], ()))
                    for c in ACTIVE_CARRIERS if c["master_col"]}

    data = {
        "warehouses": [whs[k] for k in sorted(whs)],
        "active_carriers": ACTIVE_CARRIERS,
        "price_cards": PRICE_CARDS,
        "master_col_by_id": MASTER_COL_BY_ID,
        "master_zones": master_zones,
        "billing_zone_mapping": billing,
        "actual_zone_name": actual,
        "na_sentinels": sorted(NA_SENTINELS),
    }
    with open(B2C_JSON, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"Wrote {B2C_JSON}")
    print("Done.")


if __name__ == "__main__":
    main()